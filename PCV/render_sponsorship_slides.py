import os
from PIL import Image, ImageDraw, ImageFont

# 16:9 Full HD Presentation Slide Size (1920 x 1080)
W, H = 1920, 1080
FONT_PATH = '/System/Library/Fonts/AppleSDGothicNeo.ttc'

def get_font(size, weight=16):
    # 16: Heavy, 14: ExtraBold, 6: Bold, 4: SemiBold, 0: Regular
    return ImageFont.truetype(FONT_PATH, size, index=weight)

# Color Palette
BG_DARK = (15, 17, 22)
SURFACE_CARD = (23, 26, 34)
SURFACE_CARD_A = (28, 30, 42)
BORDER_COLOR = (42, 47, 60)
BORDER_ORANGE = (255, 90, 0)
ORANGE_ACCENT = (255, 90, 0)
GOLD_ACCENT = (255, 180, 40)
TEXT_WHITE = (245, 248, 252)
TEXT_GRAY = (150, 160, 175)
TEXT_MUTED = (100, 110, 125)
GREEN_ACCENT = (46, 204, 113)

# -------------------------------------------------------------
# SLIDE 1: 비교표 및 핵심 요약 (Comparison Table & Key Summary)
# -------------------------------------------------------------
def render_slide1():
    img = Image.new('RGB', (W, H), BG_DARK)
    draw = ImageDraw.Draw(img)
    
    # Top Tag & Header
    draw.rectangle([(80, 50), (250, 80)], fill=(40, 20, 15), outline=ORANGE_ACCENT, width=1)
    draw.text((94, 55), "BHL 2026 SPONSORSHIP", font=get_font(13, 14), fill=ORANGE_ACCENT)
    
    draw.text((80, 92), "BHL 2026 앰배서더 연계 공식 협찬 패키지", font=get_font(38, 16), fill=TEXT_WHITE)
    draw.text((80, 144), "7.4만 인스타툰 작가 꿀차(@ggul.cha)와 함께하는 2030 여성 타깃 집중 마케팅 제안", font=get_font(18, 4), fill=TEXT_GRAY)

    # 4 Meta Badges at top
    badges = [
        ("공식 앰배서더", "꿀차 작가 (@ggul.cha)", "7.4만 팔로워 인스타툰 작가"),
        ("핵심 타깃층", "2030 여성 80% 이상", "대한민국 거주 / 강력한 팬덤"),
        ("채널 파워", "툰 10만~15만 / 릴스 5~20만", "압도적 도달 및 저장/공유율"),
        ("타깃 핏 (Fit)", "운동 · 헬스 · 웰니스 · 뷰티", "여성 라이프스타일에 최적화")
    ]
    bw = 415
    for i, (title, main, sub) in enumerate(badges):
        bx = 80 + i * (bw + 23)
        by = 188
        draw.rectangle([(bx, by), (bx + bw, by + 95)], fill=SURFACE_CARD, outline=BORDER_COLOR, width=1)
        # Accent indicator
        draw.rectangle([(bx, by), (bx + 4, by + 95)], fill=ORANGE_ACCENT if i == 0 else (60, 68, 85))
        draw.text((bx + 18, by + 14), title, font=get_font(13, 6), fill=ORANGE_ACCENT if i == 0 else TEXT_MUTED)
        draw.text((bx + 18, by + 36), main, font=get_font(18, 14), fill=TEXT_WHITE)
        draw.text((bx + 18, by + 66), sub, font=get_font(13, 0), fill=TEXT_GRAY)

    # Table Grid Columns
    # Col 1: 구분 (320px)  (80 to 400)
    # Col 2: GRADE A (740px) (400 to 1140)
    # Col 3: GRADE B (700px) (1140 to 1840)
    ty = 312
    tw = 1760
    header_h = 75
    
    rows = [
        ("캐러셀 (인스타툰)",
         "스토리텔링형 캐러셀 10장 (단독 편성)\n• 최근 조회수 10만~15만 회 / 제품 중심 맞춤 에피소드 연출",
         "- (미포함)",
         90),
        ("현장 릴스 (15초)",
         "대회 현장 스케치 및 제품 노출 릴스 1편\n• 최근 평균 조회수 5만~20만 회 / 선수들의 실제 제품 사용 장면 노출",
         "대회 현장 스케치 및 제품 노출 릴스 1편\n• 최근 평균 조회수 5만~20만 회 / 선수들의 실제 제품 사용 장면 노출",
         90),
        ("블로그 포스팅",
         "대회 스케치 및 제품 상세 포스팅 (1편)\n• 네이버 키워드 상위 SEO 노출, 상세 리뷰 및 구매 링크 연계",
         "대회 스케치 및 제품 상세 포스팅 (1편)\n• 네이버 키워드 상위 SEO 노출, 상세 리뷰 및 구매 링크 연계",
         90),
        ("결과분석리포트",
         "D+14 공식 성과 분석 결과보고서 제공 (PDF)\n• 도달수, 조회수, 저장/공유수, 댓글 반응 데이터 공식 취합 전달",
         "D+14 공식 성과 분석 결과보고서 제공 (PDF)\n• 도달수, 조회수, 저장/공유수, 댓글 반응 데이터 공식 취합 전달",
         90),
        ("정상 광고비",
         "2,200,000원",
         "700,000원",
         65),
        ("BHL 특별 협찬가",
         "1,200,000원  (VAT 별도 / 약 45% 특별 할인)",
         "300,000원  (VAT 별도 / 약 57% 특별 할인)",
         110)
    ]
    
    total_table_h = header_h + sum(r[3] for r in rows)
    
    # Outer Frame
    draw.rectangle([(80, ty), (80 + tw, ty + total_table_h)], fill=SURFACE_CARD, outline=BORDER_COLOR, width=1)
    
    # Column Header Backgrounds
    draw.rectangle([(80, ty), (400, ty + header_h)], fill=(28, 32, 42))
    draw.rectangle([(400, ty), (1140, ty + header_h)], fill=(45, 25, 20)) # Grade A Header (warm orange tint)
    draw.rectangle([(1140, ty), (1840, ty + header_h)], fill=(28, 32, 42))
    
    # Grade A Column Highlight Outline
    draw.rectangle([(400, ty), (1140, ty + total_table_h)], outline=ORANGE_ACCENT, width=2)
    
    # Header Titles
    draw.text((105, ty + 26), "제공 항목", font=get_font(20, 14), fill=TEXT_WHITE)
    
    # Grade A Header Text
    draw.text((430, ty + 16), "GRADE A (메인 단독 스폰서)", font=get_font(22, 16), fill=TEXT_WHITE)
    draw.rectangle([(850, ty + 20), (1110, ty + 52)], fill=ORANGE_ACCENT)
    draw.text((865, ty + 24), "RECOMMENDED / 단독 편성", font=get_font(12, 16), fill=(0, 0, 0))
    
    # Grade B Header Text
    draw.text((1170, ty + 16), "GRADE B (숏폼 & 블로그 스폰서)", font=get_font(22, 16), fill=TEXT_WHITE)
    draw.rectangle([(1590, ty + 20), (1810, ty + 52)], fill=(45, 52, 68))
    draw.text((1605, ty + 24), "가성비 바이럴 패키지", font=get_font(12, 14), fill=TEXT_GRAY)

    curr_y = ty + header_h
    for idx, (label, val_a, val_b, rh) in enumerate(rows):
        is_last = (idx == len(rows) - 1)
        is_normal_price = (idx == len(rows) - 2)
        
        # Row dividing line
        draw.line([(80, curr_y), (80 + tw, curr_y)], fill=BORDER_COLOR, width=1)
        
        if is_last:
            # Highlight special price row
            draw.rectangle([(80, curr_y), (400, curr_y + rh)], fill=(32, 22, 20))
            draw.rectangle([(400, curr_y), (1140, curr_y + rh)], fill=(55, 25, 18))
            draw.rectangle([(1140, curr_y), (1840, curr_y + rh)], fill=(28, 34, 45))
        
        # Col 1: Label
        lbl_font = get_font(18, 16 if is_last else 6)
        lbl_color = ORANGE_ACCENT if is_last else TEXT_WHITE
        draw.text((105, curr_y + (rh // 2) - 12), label, font=lbl_font, fill=lbl_color)
        
        # Col 2: Grade A Value
        if is_last:
            draw.text((430, curr_y + 20), "1,200,000원", font=get_font(32, 16), fill=ORANGE_ACCENT)
            draw.text((650, curr_y + 30), "(VAT 별도 / 약 45% 특별 할인)", font=get_font(16, 6), fill=GOLD_ACCENT)
            draw.text((430, curr_y + 68), "• 정상가 220만 원 대비 100만 원 할인 혜택  • 인스타툰 10장 단독 브랜디드 스토리텔링", font=get_font(14, 4), fill=TEXT_GRAY)
        elif is_normal_price:
            draw.text((430, curr_y + (rh // 2) - 10), val_a, font=get_font(18, 4), fill=TEXT_MUTED)
        else:
            lines = val_a.split('\n')
            draw.text((430, curr_y + 18), lines[0], font=get_font(17, 14), fill=TEXT_WHITE)
            if len(lines) > 1:
                draw.text((430, curr_y + 48), lines[1], font=get_font(14, 0), fill=TEXT_GRAY)
                
        # Col 3: Grade B Value
        if is_last:
            draw.text((1170, curr_y + 20), "300,000원", font=get_font(32, 16), fill=TEXT_WHITE)
            draw.text((1355, curr_y + 30), "(VAT 별도 / 약 57% 파격 할인)", font=get_font(16, 6), fill=GREEN_ACCENT)
            draw.text((1170, curr_y + 68), "• 정상가 70만 원 대비 40만 원 할인 혜택  • 릴스 15초 + 블로그 SEO + 결과보고서", font=get_font(14, 4), fill=TEXT_GRAY)
        elif is_normal_price:
            draw.text((1170, curr_y + (rh // 2) - 10), val_b, font=get_font(18, 4), fill=TEXT_MUTED)
        else:
            if val_b == "- (미포함)":
                draw.text((1170, curr_y + (rh // 2) - 10), val_b, font=get_font(17, 4), fill=TEXT_MUTED)
            else:
                lines = val_b.split('\n')
                draw.text((1170, curr_y + 18), lines[0], font=get_font(17, 14), fill=TEXT_WHITE)
                if len(lines) > 1:
                    draw.text((1170, curr_y + 48), lines[1], font=get_font(14, 0), fill=TEXT_GRAY)
                    
        curr_y += rh

    # Bottom notes
    draw.text((80, ty + total_table_h + 20), "* 모든 패키지는 참가 선수 및 운영진 전원(약 120명) 웰컴 기프트백 제품 동봉 샘플링을 기본 지원합니다.", font=get_font(14, 4), fill=TEXT_MUTED)
    draw.text((1680, ty + total_table_h + 20), "SLIDE 01 / 02", font=get_font(14, 14), fill=ORANGE_ACCENT)
    
    return img


# -------------------------------------------------------------
# SLIDE 2: 세부 구성 안내 & 실행 프로세스 (Detailed Specs & Workflow)
# -------------------------------------------------------------
def render_slide2():
    img = Image.new('RGB', (W, H), BG_DARK)
    draw = ImageDraw.Draw(img)
    
    # Top Tag & Header
    draw.rectangle([(80, 50), (220, 80)], fill=(40, 20, 15), outline=ORANGE_ACCENT, width=1)
    draw.text((94, 55), "DETAILED SPECS", font=get_font(13, 14), fill=ORANGE_ACCENT)
    
    draw.text((80, 92), "협찬 등급별 세부 구성 & 사후 성과 관리 프로세스", font=get_font(38, 16), fill=TEXT_WHITE)
    draw.text((80, 144), "단독 스토리텔링부터 숏폼 바이럴, 네이버 검색 SEO, D+14 데이터 리포트까지 체계적인 실행을 보장합니다.", font=get_font(18, 4), fill=TEXT_GRAY)

    card_y = 195
    card_h = 595
    card_w = 860
    
    # ------------------ LEFT CARD: GRADE A ------------------
    x_a = 80
    draw.rectangle([(x_a, card_y), (x_a + card_w, card_y + card_h)], fill=SURFACE_CARD_A, outline=BORDER_ORANGE, width=2)
    
    # Header bar
    draw.rectangle([(x_a, card_y), (x_a + card_w, card_y + 75)], fill=(38, 26, 28))
    draw.text((x_a + 25, card_y + 16), "GRADE A  |  메인 단독 스폰서", font=get_font(24, 16), fill=TEXT_WHITE)
    draw.text((x_a + card_w - 240, card_y + 18), "120만 원 (VAT 별도)", font=get_font(22, 16), fill=ORANGE_ACCENT)
    draw.text((x_a + 25, card_y + 48), "단독 에피소드 캐러셀 10장 + 릴스 15초 + 블로그 SEO + 결과분석리포트", font=get_font(14, 4), fill=GOLD_ACCENT)
    
    items_a = [
        ("01", "단독 홍보툰 (10장 캐러셀)", "정상가 150만 원",
         "• 브랜드 특장점과 메시지가 웹툰 에피소드 안에 자연스럽게 녹아드는 단독 스토리텔링\n• 꿀차 작가 채널 최근 평균 조회수 10만~15만 회 보장형 도달\n• 2030 여성 타깃의 폭발적인 공감과 브랜드 호감도 형성"),
        ("02", "현장 스케치 릴스 (15초)", "정상가 60만 원",
         "• 대회 현장 선수들의 실제 제품 음용 / 착용 / 사용 장면을 15초 숏폼으로 제작\n• 최근 릴스 평균 조회수 5만~20만 회 알고리즘 바이럴 노출\n• 박진감 넘치는 현장감과 자연스러운 PPL 결합"),
        ("03", "블로그 상세 포스팅 (SEO)", "정상가 10만 원",
         "• 네이버 상위 노출 키워드 타깃팅, 정밀한 제품 리뷰 및 특장점 소개\n• 공식 자사몰 및 스마트스토어 구매 링크 직접 연결로 실구매 전환 유도"),
        ("04", "D+14 공식 성과 분석 보고서 (PDF)", "기본 포함",
         "• 발행 후 2주간의 조회수, 도달수, 저장수, 공유수, 댓글 반응 데이터를 취합\n• 기업 내부 결재 및 마케팅 성과 증빙이 가능한 완성형 PDF 리포트 제공")
    ]
    
    iy = card_y + 95
    for num, title, price_tag, desc in items_a:
        draw.rectangle([(x_a + 25, iy), (x_a + 55, iy + 26)], fill=ORANGE_ACCENT)
        draw.text((x_a + 30, iy + 4), num, font=get_font(13, 16), fill=(0, 0, 0))
        draw.text((x_a + 65, iy + 2), title, font=get_font(18, 14), fill=TEXT_WHITE)
        draw.text((x_a + card_w - 170, iy + 4), price_tag, font=get_font(13, 6), fill=GOLD_ACCENT)
        
        dy = iy + 30
        for line in desc.split('\n'):
            draw.text((x_a + 65, dy), line, font=get_font(13, 0), fill=TEXT_GRAY)
            dy += 20
        iy += 120

    # ------------------ RIGHT CARD: GRADE B ------------------
    x_b = 980
    draw.rectangle([(x_b, card_y), (x_b + card_w, card_y + card_h)], fill=SURFACE_CARD, outline=BORDER_COLOR, width=1)
    
    # Header bar
    draw.rectangle([(x_b, card_y), (x_b + card_w, card_y + 75)], fill=(28, 32, 42))
    draw.text((x_b + 25, card_y + 16), "GRADE B  |  숏폼 & 블로그 스폰서", font=get_font(24, 16), fill=TEXT_WHITE)
    draw.text((x_b + card_w - 230, card_y + 18), "30만 원 (VAT 별도)", font=get_font(22, 16), fill=GREEN_ACCENT)
    draw.text((x_b + 25, card_y + 48), "현장 릴스 15초 + 블로그 상세 포스팅 + D+14 결과분석리포트", font=get_font(14, 4), fill=TEXT_GRAY)
    
    items_b = [
        ("01", "현장 스케치 릴스 (15초 PPL)", "정상가 60만 원",
         "• 대회 현장 공식 스케치 릴스 영상 내 제품 음용/사용 장면 1회 자연 노출\n• 최근 릴스 평균 조회수 5만~20만 회 알고리즘 바이럴 혜택 공유\n• 30만 원 소액 예산으로 인스타그램 숏폼 대규모 도달 확보"),
        ("02", "블로그 상세 포스팅 (SEO)", "정상가 10만 원",
         "• 대회 종합 스케치 및 현장 후기 포스팅 내 브랜드/제품 전용 소개 섹션\n• 고화질 현장 사진 3장 이상 첨부 및 브랜드 공식 링크 삽입\n• 네이버 포털 검색 시 영구적인 콘텐츠 아카이빙 효과"),
        ("03", "D+14 공식 성과 분석 보고서 (PDF)", "기본 포함",
         "• 릴스 및 블로그 콘텐츠의 실제 도달수와 반응 데이터를 정리한 공식 리포트\n• 소액 협찬이라도 마케팅 성과를 투명하게 수치화하여 증빙 제공"),
        ("04", "현장 샘플링 연계 지원", "기본 포함",
         "• 대회 참가 선수 및 운영진 전원(약 120명) 웰컴 기프트백 제품 동봉\n• 현장에서 2030 여성 핵심 타깃에게 직접 제품을 쥐여주는 강력한 체험 마케팅")
    ]
    
    iy = card_y + 95
    for num, title, price_tag, desc in items_b:
        draw.rectangle([(x_b + 25, iy), (x_b + 55, iy + 26)], fill=(50, 60, 80))
        draw.text((x_b + 30, iy + 4), num, font=get_font(13, 16), fill=TEXT_WHITE)
        draw.text((x_b + 65, iy + 2), title, font=get_font(18, 14), fill=TEXT_WHITE)
        draw.text((x_b + card_w - 170, iy + 4), price_tag, font=get_font(13, 6), fill=GREEN_ACCENT)
        
        dy = iy + 30
        for line in desc.split('\n'):
            draw.text((x_b + 65, dy), line, font=get_font(13, 0), fill=TEXT_GRAY)
            dy += 20
        iy += 120

    # ------------------ BOTTOM WORKFLOW BAR ------------------
    by = 820
    bw = 1760
    bh = 195
    draw.rectangle([(80, by), (80 + bw, by + bh)], fill=SURFACE_CARD, outline=BORDER_COLOR, width=1)
    
    draw.text((105, by + 18), "진행 프로세스 & 사후 성과 보증 (Timeline)", font=get_font(18, 16), fill=TEXT_WHITE)
    draw.text((450, by + 20), "사전 기획부터 사후 보고서까지 원스톱으로 체계적인 커뮤니케이션을 진행합니다.", font=get_font(14, 0), fill=TEXT_GRAY)
    
    steps = [
        ("STEP 01", "협찬 확정 & 제품 전달", "D-20", "협찬 등급 확정 및 기프트백용 샘플 제품 수령"),
        ("STEP 02", "콘티 검수 & 기획 회의", "D-10", "인스타툰 및 릴스 핵심 노출 포인트 사전 조율"),
        ("STEP 03", "대회 당일 현장 실행", "D-DAY", "웰컴백 배포, 선수 실사용, 현장 영상 촬영"),
        ("STEP 04", "콘텐츠 순차 발행", "D+7", "꿀차 계정 인스타툰, 릴스, 네이버 블로그 업로드"),
        ("STEP 05", "D+14 성과 분석 리포트", "D+14", "최종 도달수, 조회수, 반응 데이터 취합 PDF 전달")
    ]
    
    sw = (bw - 50) // 5
    for i, (step_tag, step_title, step_day, step_desc) in enumerate(steps):
        sx = 105 + i * sw
        s_box_y = by + 56
        draw.rectangle([(sx, s_box_y), (sx + sw - 15, s_box_y + 115)], fill=(18, 20, 26), outline=BORDER_COLOR, width=1)
        
        # Step tag
        draw.text((sx + 14, s_box_y + 12), step_tag, font=get_font(12, 14), fill=ORANGE_ACCENT if i in [0, 4] else TEXT_MUTED)
        draw.rectangle([(sx + sw - 68, s_box_y + 10), (sx + sw - 22, s_box_y + 28)], fill=(32, 38, 50))
        draw.text((sx + sw - 62, s_box_y + 12), step_day, font=get_font(11, 14), fill=GOLD_ACCENT)
        
        draw.text((sx + 14, s_box_y + 36), step_title, font=get_font(15, 14), fill=TEXT_WHITE)
        
        f_desc = get_font(12, 0)
        words = step_desc.split(' ')
        cur = ""
        lines = []
        for w in words:
            if len(cur + " " + w) < 18:
                cur += (" " if cur else "") + w
            else:
                lines.append(cur)
                cur = w
        if cur: lines.append(cur)
        
        sdy = s_box_y + 64
        for l in lines[:2]:
            draw.text((sx + 14, sdy), l, font=f_desc, fill=TEXT_GRAY)
            sdy += 18

    draw.text((1680, by + bh + 18), "SLIDE 02 / 02", font=get_font(14, 14), fill=ORANGE_ACCENT)
    return img

print("Rendering Slide 1...")
s1 = render_slide1()
s1_path = '/Users/hnky/agyspace/PCV/assets/bhl_sponsorship_slide1.png'
s1.save(s1_path, quality=95)

print("Rendering Slide 2...")
s2 = render_slide2()
s2_path = '/Users/hnky/agyspace/PCV/assets/bhl_sponsorship_slide2.png'
s2.save(s2_path, quality=95)

# Combined PDF
pdf_path = '/Users/hnky/agyspace/PCV/BHL_2026_SPONSORSHIP_PACKAGE.pdf'
s1.save(pdf_path, "PDF", resolution=150.0, save_all=True, append_images=[s2])
print("PDF created.")

# PowerPoint (.pptx) creation
from pptx import Presentation
from pptx.util import Inches

prs = Presentation()
# 16:9 widescreen dimensions in inches (13.333 x 7.5)
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_slide_layout = prs.slide_layouts[6]

slide1 = prs.slides.add_slide(blank_slide_layout)
slide1.shapes.add_picture(s1_path, Inches(0), Inches(0), Inches(13.333), Inches(7.5))

slide2 = prs.slides.add_slide(blank_slide_layout)
slide2.shapes.add_picture(s2_path, Inches(0), Inches(0), Inches(13.333), Inches(7.5))

pptx_path = '/Users/hnky/agyspace/PCV/BHL_2026_SPONSORSHIP_PACKAGE.pptx'
prs.save(pptx_path)
print(f"PPTX created successfully at: {pptx_path}")

