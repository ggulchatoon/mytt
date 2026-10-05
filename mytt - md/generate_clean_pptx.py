import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Design system from original PPT
FONT_TITLE = 'Gowun Dodum'
FONT_BODY = 'Noto Sans KR'

COLOR_BROWN_DEEP = RGBColor(0x3E, 0x2C, 0x21)   # #3E2C21 (제목 짙은 브라운)
COLOR_BROWN_TEXT = RGBColor(0x5C, 0x4A, 0x3E)   # #5C4A3E (본문 텍스트)
COLOR_BROWN_MUTED = RGBColor(0x8C, 0x6A, 0x54)  # #8C6A54 (설명/캡션)
COLOR_ORANGE = RGBColor(0xE7, 0x6F, 0x51)       # #E76F51 (원본 오렌지 포인트)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]
bg_img = '/Users/hnky/agyspace/slide_bg.png'

def create_slide():
    slide = prs.slides.add_slide(blank_layout)
    if os.path.exists(bg_img):
        slide.shapes.add_picture(bg_img, 0, 0, prs.slide_width, prs.slide_height)
    return slide

def add_headline(slide, text):
    """우상단/좌상단 헤드라인: 오렌지 바(0.1 x 0.31) + Gowun Dodum 27pt"""
    # Orange vertical bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.625), Inches(0.674), Inches(0.104), Inches(0.313))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_ORANGE
    bar.line.fill.background()

    # Headline text
    tbox = slide.shapes.add_textbox(Inches(0.867), Inches(0.625), Inches(12.0), Inches(0.454))
    tf = tbox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = FONT_TITLE
    p.font.size = Pt(27.0)
    p.font.bold = True
    p.font.color.rgb = COLOR_BROWN_DEEP

def add_center_title(slide, main_text, sub_text=None, top_inch=2.7, main_size=40.5):
    """중앙 정렬 타이틀 슬라이드 (지난 PPT와 동일한 규격)"""
    tbox = slide.shapes.add_textbox(Inches(1.5), Inches(top_inch), Inches(10.333), Inches(3.0))
    tf = tbox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]
    p.text = main_text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_TITLE
    p.font.size = Pt(main_size)
    p.font.bold = True
    p.font.color.rgb = COLOR_BROWN_DEEP
    
    if sub_text:
        p.space_after = Pt(20)
        p_sub = tf.add_paragraph()
        p_sub.text = sub_text
        p_sub.alignment = PP_ALIGN.CENTER
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(17.0)
        p_sub.font.color.rgb = COLOR_BROWN_TEXT

# ==================== SLIDE 1 ====================
s1 = create_slide()
# Center layout exactly matching original slide 1
tbox = s1.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.333), Inches(1.2))
tf = tbox.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

p0 = tf.paragraphs[0]
p0.text = "mytt 2회차 파일럿 클래스"
p0.alignment = PP_ALIGN.CENTER
p0.font.name = FONT_BODY
p0.font.size = Pt(16.0)
p0.font.bold = True
p0.font.color.rgb = COLOR_ORANGE
p0.space_after = Pt(12)

p1 = tf.add_paragraph()
p1.text = "“안녕하세요”"
p1.alignment = PP_ALIGN.CENTER
p1.font.name = FONT_TITLE
p1.font.size = Pt(43.5)
p1.font.bold = True
p1.font.color.rgb = COLOR_BROWN_DEEP

# Body text
bbox = s1.shapes.add_textbox(Inches(2.0), Inches(3.9), Inches(9.333), Inches(2.2))
btf = bbox.text_frame
btf.word_wrap = True
btf.margin_left = btf.margin_top = btf.margin_right = btf.margin_bottom = 0

p2 = btf.paragraphs[0]
p2.text = "채팅창에 반갑게 인사 나누어 보아요 🌿"
p2.alignment = PP_ALIGN.CENTER
p2.font.name = FONT_BODY
p2.font.size = Pt(17.0)
p2.font.color.rgb = COLOR_BROWN_TEXT
p2.space_after = Pt(24)

p3 = btf.add_paragraph()
p3.text = "소리는 OFF / 화면은 ON 해주시면 진심으로 감사드리겠습니다… (필수 x 자유 o)"
p3.alignment = PP_ALIGN.CENTER
p3.font.name = FONT_BODY
p3.font.size = Pt(13.5)
p3.font.color.rgb = COLOR_BROWN_MUTED
p3.space_after = Pt(10)

p4 = btf.add_paragraph()
p4.text = "Ai를 준비해주세요! (클로드, 제미나이, 챗GPT 등)"
p4.alignment = PP_ALIGN.CENTER
p4.font.name = FONT_BODY
p4.font.size = Pt(14.5)
p4.font.bold = True
p4.font.color.rgb = COLOR_ORANGE

# ==================== SLIDE 2 ====================
s2 = create_slide()
add_headline(s2, "지난 시간 복습")

# 2 Columns for review
col1 = s2.shapes.add_textbox(Inches(1.0), Inches(2.3), Inches(5.3), Inches(4.5))
tf1 = col1.text_frame
tf1.word_wrap = True
tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

p = tf1.paragraphs[0]
p.text = "시간 영토"
p.font.name = FONT_TITLE
p.font.size = Pt(24.0)
p.font.bold = True
p.font.color.rgb = COLOR_ORANGE
p.space_after = Pt(16)

items_territory = [
    ("수면", "바다"),
    ("고정스케줄", "언덕"),
    ("휴식", "평원"),
    ("일", "나의 시간농장")
]
for name, meta in items_territory:
    pi = tf1.add_paragraph()
    pi.text = f"• {name}  ({meta})"
    pi.font.name = FONT_BODY
    pi.font.size = Pt(16.0)
    pi.font.color.rgb = COLOR_BROWN_TEXT
    pi.space_after = Pt(12)

col2 = s2.shapes.add_textbox(Inches(6.8), Inches(2.3), Inches(5.5), Inches(4.5))
tf2 = col2.text_frame
tf2.word_wrap = True
tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0

p = tf2.paragraphs[0]
p.text = "딥워크 (Deep Work)"
p.font.name = FONT_TITLE
p.font.size = Pt(24.0)
p.font.bold = True
p.font.color.rgb = COLOR_ORANGE
p.space_after = Pt(16)

p_sub = tf2.add_paragraph()
p_sub.text = "• 최대 4시간"
p_sub.font.name = FONT_BODY
p_sub.font.size = Pt(18.0)
p_sub.font.bold = True
p_sub.font.color.rgb = COLOR_BROWN_DEEP
p_sub.space_after = Pt(12)

p_desc = tf2.add_paragraph()
p_desc.text = "창작자, 연구가가 몰입해 결과물을 내는 시간"
p_desc.font.name = FONT_BODY
p_desc.font.size = Pt(16.0)
p_desc.font.color.rgb = COLOR_BROWN_TEXT

# ==================== SLIDE 3 ====================
s3 = create_slide()
add_headline(s3, "1회차 피드백")

col1 = s3.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.4), Inches(4.5))
tf1 = col1.text_frame
tf1.word_wrap = True
tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

p = tf1.paragraphs[0]
p.text = "공통 피드백"
p.font.name = FONT_TITLE
p.font.size = Pt(21.0)
p.font.bold = True
p.font.color.rgb = COLOR_ORANGE
p.space_after = Pt(16)

common_fbs = [
    "• 워크북에 입력할 내용이 많아 조금 부담스러웠다",
    "• 소감 나눔 시간이 조금 길었다",
    "• 워크북(이론)과 이후 활동(바이브 코딩)이 앞으로 유기적으로 연결되면 좋겠다"
]
for item in common_fbs:
    pi = tf1.add_paragraph()
    pi.text = item
    pi.font.name = FONT_BODY
    pi.font.size = Pt(15.0)
    pi.font.color.rgb = COLOR_BROWN_TEXT
    pi.space_after = Pt(12)

col2 = s3.shapes.add_textbox(Inches(6.9), Inches(2.2), Inches(5.4), Inches(4.5))
tf2 = col2.text_frame
tf2.word_wrap = True
tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0

p = tf2.paragraphs[0]
p.text = "기타 피드백"
p.font.name = FONT_TITLE
p.font.size = Pt(21.0)
p.font.bold = True
p.font.color.rgb = COLOR_ORANGE
p.space_after = Pt(16)

etc_fbs = [
    "• 나의 상황에 맞지 않는 부분이 있어 고민되었다",
    "• 주말 루틴도 추가해보고 싶다 (업데이트 예정)"
]
for item in etc_fbs:
    pi = tf2.add_paragraph()
    pi.text = item
    pi.font.name = FONT_BODY
    pi.font.size = Pt(15.0)
    pi.font.color.rgb = COLOR_BROWN_TEXT
    pi.space_after = Pt(12)

p_end = tf2.add_paragraph()
p_end.space_before = Pt(16)
p_end.text = "-> 더 귀기울여 듣고 발전시켜보겠습니다!"
p_end.font.name = FONT_BODY
p_end.font.size = Pt(15.0)
p_end.font.bold = True
p_end.font.color.rgb = COLOR_BROWN_DEEP

# ==================== SLIDE 4 ====================
s4 = create_slide()
add_headline(s4, "<나눔>")

box = s4.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(5.0))
tf = box.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

p = tf.paragraphs[0]
p.text = "Q. 지난 시간 이후, 어떠셨나요?"
p.font.name = FONT_TITLE
p.font.size = Pt(25.5)
p.font.bold = True
p.font.color.rgb = COLOR_ORANGE
p.space_after = Pt(16)

items_s4 = [
    "• 일상에서 나의 시간 사용 / 시간 도둑 관찰",
    "• 어렵고 힘들었던 점",
    "등등…"
]
for it in items_s4:
    pi = tf.add_paragraph()
    pi.text = it
    pi.font.name = FONT_BODY
    pi.font.size = Pt(15.5)
    pi.font.color.rgb = COLOR_BROWN_TEXT
    pi.space_after = Pt(8)

p_ex_head = tf.add_paragraph()
p_ex_head.space_before = Pt(16)
p_ex_head.text = "ex- 저는 하루 4시간 딥워크 시간을 시도해보았는데,\n     마음 먹으니 2-3시간정도 가능하더라고요!"
p_ex_head.font.name = FONT_BODY
p_ex_head.font.size = Pt(14.5)
p_ex_head.font.color.rgb = COLOR_BROWN_MUTED
p_ex_head.space_after = Pt(10)

p_ex2 = tf.add_paragraph()
p_ex2.text = "ex- 저는 시간 도둑이 너무 많아 작업 시작이 어렵다는 걸 깨달았어요."
p_ex2.font.name = FONT_BODY
p_ex2.font.size = Pt(14.5)
p_ex2.font.color.rgb = COLOR_BROWN_MUTED
p_ex2.space_after = Pt(16)

p_action = tf.add_paragraph()
p_action.text = "-> 채팅창에 경험하고 느낀 점을 한 줄 나누어보세요"
p_action.font.name = FONT_BODY
p_action.font.size = Pt(16.0)
p_action.font.bold = True
p_action.font.color.rgb = COLOR_BROWN_DEEP

# ==================== SLIDE 5 ====================
s5 = create_slide()
add_center_title(s5, "원래 저희의 의도", "-> 으쌰으쌰 인증샷 올리고 다같이 자신만의 루틴을 만들어가는 모임\n-> ???\n-> 제가 더 노력하겠습니다", top_inch=2.3, main_size=36.0)

# ==================== SLIDE 6 ====================
s6 = create_slide()
add_headline(s6, "<나눔>")

box = s6.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(5.0))
tf = box.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

p = tf.paragraphs[0]
p.text = "나의 Ai 사용 경험"
p.font.name = FONT_TITLE
p.font.size = Pt(25.5)
p.font.bold = True
p.font.color.rgb = COLOR_ORANGE
p.space_after = Pt(16)

items_s6 = [
    "- 언제 무슨 용도로 사용해보았는지",
    "- 웹사이트 (html) 를 만들어본 적 있는지"
]
for it in items_s6:
    pi = tf.add_paragraph()
    pi.text = it
    pi.font.name = FONT_BODY
    pi.font.size = Pt(16.0)
    pi.font.color.rgb = COLOR_BROWN_TEXT
    pi.space_after = Pt(10)

p_ex1 = tf.add_paragraph()
p_ex1.space_before = Pt(20)
p_ex1.text = "ex- 그림 그려봐! 이렇게 갖고 놀아본 (?) 정도에요 ㅎㅎ"
p_ex1.font.name = FONT_BODY
p_ex1.font.size = Pt(15.0)
p_ex1.font.color.rgb = COLOR_BROWN_MUTED
p_ex1.space_after = Pt(10)

p_ex2 = tf.add_paragraph()
p_ex2.text = "ex- 여행 갈 때 일정표를 html로 만들어 본 적 있습니다!"
p_ex2.font.name = FONT_BODY
p_ex2.font.size = Pt(15.0)
p_ex2.font.color.rgb = COLOR_BROWN_MUTED

# ==================== SLIDE 7 ====================
s7 = create_slide()
add_center_title(s7, "걱정마세요!", "차근차근 함께 해볼게요", top_inch=2.8, main_size=43.5)

# ==================== SLIDE 8 ====================
s8 = create_slide()
add_center_title(s8, "예제 : 나의 몰입 시간을 지켜줄\n[딥워크 타이머 만들기]", None, top_inch=2.8, main_size=36.0)

# ==================== SLIDE 9 ====================
s9 = create_slide()
add_center_title(s9, "얼마나 걸렸을까요?", "(이미지 : 클로드를 활용하는 깨굴이 옆에서 구경하는 꿀차)", top_inch=2.7, main_size=40.5)

# ==================== SLIDE 10 ====================
s10 = create_slide()
add_center_title(s10, "단 5분이면 충분합니다", "(이미지 : 깨굴이 옆의 놀래는 꿀차)", top_inch=2.7, main_size=40.5)

# ==================== SLIDE 11 ====================
s11 = create_slide()
add_center_title(s11, "내가 만들고 싶은 것의\n‘의도’와 ‘기능’을 잘 전달해주었기 때문에", "-> Ai와 3번 정도 대화하니 완벽하게 만들어졌다\n(사실 5분은 뻥이고 수정까지 10분)", top_inch=2.3, main_size=32.0)

# ==================== SLIDE 12 ====================
s12 = create_slide()
add_center_title(s12, "당황하지 마세요", "처음 해보면 생각과 다를 수 있어요\n\n하지만 고쳐나가면 됩니다", top_inch=2.4, main_size=40.5)

# ==================== SLIDE 13 (스크립트 14번) ====================
s13 = create_slide()
add_headline(s13, "바이브 코딩의 3가지 요소")

box = s13.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(4.5))
tf = box.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

items_s14 = [
    ("기획", "내가 만들고 싶은 것은? (의도와 쓰임새, 용어)"),
    ("기능", "작동 규칙 (클릭, 타이머, 저장 기능 등)"),
    ("디자인", "보면 기분이 좋거든요 (색감, 무드 등)")
]
for idx, (title, desc) in enumerate(items_s14):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    p.text = f"{title} - {desc}"
    p.font.name = FONT_BODY
    p.font.size = Pt(18.0)
    p.font.bold = True if idx == 0 else False
    p.font.color.rgb = COLOR_BROWN_DEEP
    p.space_after = Pt(20)

# ==================== SLIDE 14 (스크립트 15번) ====================
s14 = create_slide()
add_headline(s14, "2회차 모임에 집중할 것")

box = s14.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(4.5))
tf = box.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

items_s15 = [
    "기획 - 내가 쓰고 싶은 것은? (의도과 쓰임새, 용어)  ✅",
    "기능 - 작동 규칙 (클릭, 타이머, 저장기능 등)  ✅",
    "디자인 - 보면 기분이 좋거든요 (색감, 무드 등)"
]
for idx, text in enumerate(items_s15):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    p.text = text
    p.font.name = FONT_BODY
    p.font.size = Pt(18.0)
    p.font.bold = True if '✅' in text else False
    p.font.color.rgb = COLOR_ORANGE if '✅' in text else COLOR_BROWN_MUTED
    p.space_after = Pt(20)

# ==================== SLIDE 15 (스크립트 16번) ====================
s15 = create_slide()
add_headline(s15, "[기획하기]")

box = s15.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.8))
tf = box.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

p = tf.paragraphs[0]
p.text = "예제 : 나의 몰입 시간을 지켜줄\n[딥워크 타이머 만들기]"
p.font.name = FONT_TITLE
p.font.size = Pt(26.0)
p.font.bold = True
p.font.color.rgb = COLOR_BROWN_DEEP
p.space_after = Pt(24)

p2 = tf.add_paragraph()
p2.text = "내가 쓰고 싶은 것은?"
p2.font.name = FONT_BODY
p2.font.size = Pt(18.0)
p2.font.bold = True
p2.font.color.rgb = COLOR_ORANGE
p2.space_after = Pt(12)

p3 = tf.add_paragraph()
p3.text = "* 나의 집중 시간을 도와줄 딥워크 타이머. 시간의 흐름에 따라, 산을 오르는 컨셉으로."
p3.font.name = FONT_BODY
p3.font.size = Pt(16.5)
p3.font.color.rgb = COLOR_BROWN_TEXT

# ==================== SLIDE 16 (스크립트 17번) ====================
s16 = create_slide()
add_headline(s16, "[기획하기]")

box = s16.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(5.2))
tf = box.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

p = tf.paragraphs[0]
p.text = "각 탭의 핵심 목적을 한줄로 적어봅시다."
p.font.name = FONT_TITLE
p.font.size = Pt(23.0)
p.font.bold = True
p.font.color.rgb = COLOR_BROWN_DEEP
p.space_after = Pt(18)

tabs17 = [
    "- 산오르기 : ",
    "- 딥워크 : ",
    "- 바구니 : ",
    "- 업무시간 : "
]
for t in tabs17:
    pi = tf.add_paragraph()
    pi.text = t
    pi.font.name = FONT_BODY
    pi.font.size = Pt(17.0)
    pi.font.color.rgb = COLOR_BROWN_TEXT
    pi.space_after = Pt(12)

p_end = tf.add_paragraph()
p_end.space_before = Pt(16)
p_end.text = "완벽한 정답은 없으니, ‘한 줄’로 정리해보고 채팅창에 올려보세요!"
p_end.font.name = FONT_BODY
p_end.font.size = Pt(15.0)
p_end.font.bold = True
p_end.font.color.rgb = COLOR_ORANGE

# ==================== SLIDE 17 (스크립트 18번) ====================
s17 = create_slide()
add_headline(s17, "[기획하기]")

box = s17.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(5.2))
tf = box.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

p = tf.paragraphs[0]
p.text = "예시"
p.font.name = FONT_TITLE
p.font.size = Pt(24.0)
p.font.bold = True
p.font.color.rgb = COLOR_ORANGE
p.space_after = Pt(16)

tabs18 = [
    ("- 산오르기 : ", "내 하루의 작업량을 측정하는 산오르기 컨셉의 타이머."),
    ("- 딥워크 : ", "내가 집중할 업무를 선택하고 산오르기에 설정할 수 있는 탭."),
    ("- 바구니 : ", "딥워크 외에, 그때그때 생각나는 일거리를 적는 체크리스트."),
    ("- 업무시간 : ", "내가 산을 얼마나 올랐는지 시간이 기록되는 탭.")
]
for label, desc in tabs18:
    pi = tf.add_paragraph()
    pi.text = label + desc
    pi.font.name = FONT_BODY
    pi.font.size = Pt(16.0)
    pi.font.color.rgb = COLOR_BROWN_TEXT
    pi.space_after = Pt(12)

# ==================== SLIDE 18 (스크립트 19번) ====================
s18 = create_slide()
add_headline(s18, "3회차 모임 미리보기")

box = s18.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.8))
tf = box.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

p1 = tf.paragraphs[0]
p1.text = "깃허브 아이디 만들기 (배포)"
p1.font.name = FONT_TITLE
p1.font.size = Pt(21.0)
p1.font.bold = True
p1.font.color.rgb = COLOR_ORANGE
p1.space_after = Pt(6)

p1_sub = tf.add_paragraph()
p1_sub.text = "- 실제로 인터넷에서 접속할 수 있게 만들어요."
p1_sub.font.name = FONT_BODY
p1_sub.font.size = Pt(16.0)
p1_sub.font.color.rgb = COLOR_BROWN_TEXT
p1_sub.space_after = Pt(24)

p2 = tf.add_paragraph()
p2.text = "내가 만든 시간관리툴 발전시키기"
p2.font.name = FONT_TITLE
p2.font.size = Pt(21.0)
p2.font.bold = True
p2.font.color.rgb = COLOR_ORANGE
p2.space_after = Pt(6)

p2_sub1 = tf.add_paragraph()
p2_sub1.text = "- 예제 : 딥워크 타이머 발전시키기"
p2_sub1.font.name = FONT_BODY
p2_sub1.font.size = Pt(16.0)
p2_sub1.font.color.rgb = COLOR_BROWN_TEXT
p2_sub1.space_after = Pt(8)

p2_sub2 = tf.add_paragraph()
p2_sub2.text = "- 나만의 아이디어로 만들어보기"
p2_sub2.font.name = FONT_BODY
p2_sub2.font.size = Pt(16.0)
p2_sub2.font.color.rgb = COLOR_BROWN_TEXT

# ==================== SLIDE 19 (스크립트 20번) ====================
s19 = create_slide()
add_headline(s19, "2회차 숙제")

box = s19.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.8))
tf = box.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

p = tf.paragraphs[0]
p.text = "이번 주 미션!"
p.font.name = FONT_TITLE
p.font.size = Pt(22.0)
p.font.bold = True
p.font.color.rgb = COLOR_ORANGE
p.space_after = Pt(16)

missions = [
    "- 나의 딥워크 시간 기록 ‘인증샷’ 올려보기 (선택)",
    "- 내가 만든 타이머 직접 사용하고 아쉬운 점 수정해보기"
]
for m in missions:
    pi = tf.add_paragraph()
    pi.text = m
    pi.font.name = FONT_BODY
    pi.font.size = Pt(16.5)
    pi.font.color.rgb = COLOR_BROWN_TEXT
    pi.space_after = Pt(12)

p_openchat = tf.add_paragraph()
p_openchat.space_before = Pt(16)
p_openchat.text = "-> 언제든 오픈채팅방에 질문하면 봐드릴게요! 자랑도 환영!"
p_openchat.font.name = FONT_BODY
p_openchat.font.size = Pt(16.0)
p_openchat.font.bold = True
p_openchat.font.color.rgb = COLOR_BROWN_DEEP

output_path = '/Users/hnky/agyspace/mytt_2회차_강의슬라이드_구원파트_v2.pptx'
prs.save(output_path)
print("Saved clean PPTX:", output_path, "Slide count:", len(prs.slides))
