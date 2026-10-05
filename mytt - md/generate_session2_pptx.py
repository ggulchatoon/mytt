import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# 1. Fonts & Colors
FONT_TITLE = 'Gowun Dodum'
FONT_BODY = 'Noto Sans KR'

COLOR_BROWN_DEEP = RGBColor(62, 44, 33)      # #3E2C21 (헤드라인 브라운)
COLOR_BROWN_TEXT = RGBColor(92, 74, 62)      # #5C4A3E (본문 텍스트)
COLOR_BROWN_MUTED = RGBColor(140, 115, 95)   # #8C735F (설명/캡션)
COLOR_ORANGE = RGBColor(240, 164, 92)        # #F0A45C (시그니처 오렌지)
COLOR_PEACH_LIGHT = RGBColor(255, 232, 214)  # #FFE8D6 (강조 배경)
COLOR_CARD_BG = RGBColor(255, 253, 248)      # #FFFDF8 (카드 배경)
COLOR_CARD_BORDER = RGBColor(237, 216, 200)  # #EDD8C8 (카드 테두리)
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_SKY_BLUE = RGBColor(137, 190, 240)
COLOR_GREEN = RGBColor(142, 196, 134)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]
bg_img = '/Users/hnky/agyspace/slide_bg.png'

def create_slide():
    slide = prs.slides.add_slide(blank_layout)
    if os.path.exists(bg_img):
        slide.shapes.add_picture(bg_img, 0, 0, prs.slide_width, prs.slide_height)
    return slide

def add_header(slide, step_label, title_text):
    if step_label:
        sbox = slide.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.333), Inches(0.35))
        tf_s = sbox.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0
        p_s = tf_s.paragraphs[0]
        p_s.text = step_label.upper()
        p_s.font.name = FONT_BODY
        p_s.font.size = Pt(13)
        p_s.font.bold = True
        p_s.font.color.rgb = COLOR_ORANGE

    if title_text:
        tbox = slide.shapes.add_textbox(Inches(1.0), Inches(0.95), Inches(11.333), Inches(0.8))
        tf_t = tbox.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = FONT_TITLE
        p_t.font.size = Pt(28)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_BROWN_DEEP

def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape

# ==================== SLIDE 1 ====================
s1 = create_slide()
# Title text box
tbox = s1.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.333), Inches(1.8))
tf = tbox.text_frame
tf.word_wrap = True
p0 = tf.paragraphs[0]
p0.text = "mytt 2회차 파일럿 클래스"
p0.alignment = PP_ALIGN.CENTER
p0.font.name = FONT_BODY
p0.font.size = Pt(16)
p0.font.bold = True
p0.font.color.rgb = COLOR_ORANGE

p1 = tf.add_paragraph()
p1.text = "“안녕하세요!” 반갑게 인사 나누어요 🌿"
p1.alignment = PP_ALIGN.CENTER
p1.font.name = FONT_TITLE
p1.font.size = Pt(36)
p1.font.bold = True
p1.font.color.rgb = COLOR_BROWN_DEEP

# Notice Card
add_card(s1, Inches(2.2), Inches(4.0), Inches(8.933), Inches(1.8), COLOR_CARD_BG, COLOR_CARD_BORDER)
nbox = s1.shapes.add_textbox(Inches(2.5), Inches(4.2), Inches(8.333), Inches(1.4))
ntf = nbox.text_frame
ntf.word_wrap = True
np1 = ntf.paragraphs[0]
np1.text = "• 소리는 OFF / 화면은 자유롭게 ON 해주시면 감사하겠습니다! (필수 X, 자유 O)"
np1.font.name = FONT_BODY
np1.font.size = Pt(16)
np1.font.color.rgb = COLOR_BROWN_TEXT

np2 = ntf.add_paragraph()
np2.text = "• 오늘 함께 실습할 AI를 미리 열어주세요! (클로드, 제미나이, 챗GPT 등)"
np2.font.name = FONT_BODY
np2.font.size = Pt(16)
np2.font.bold = True
np2.font.color.rgb = COLOR_ORANGE

# ==================== SLIDE 2 ====================
s2 = create_slide()
add_header(s2, "STEP 01 · REVIEW", "지난 시간 복습 : 시간 영토와 딥워크")

card_w = Inches(2.6)
card_h = Inches(4.3)
cards_data = [
    ("🌊 수면 · 바다", "생명의 근원이 되는 시간", "언제 자고 언제 일어나는지,\n내 하루 컨디션을 지켜주는\n가장 기초적인 시간의 바다"),
    ("⛰️ 고정스케줄 · 언덕", "일상의 뼈대를 잡는 시간", "출퇴근, 식사, 고정 미팅 등\n변동 없이 매주 콕 박혀있는\n일상의 고정된 언덕"),
    ("🌅 휴식 · 평원", "마음 편히 재충전하는 시간", "일과 후 자유롭게 쉬는 시간!\n죄책감 없이 온전히 쉬어야\n다음 날 몰입도 지켜져요"),
    ("🍎 나의 시간 농장", "진짜 결실을 맺는 시간", "바다/언덕/평원을 뺀 가변 시간!\n영감의 씨앗을 뿌리고\n나만의 결실을 거두는 비옥한 밭")
]

for idx, (title, sub, desc) in enumerate(cards_data):
    cx = Inches(1.0) + idx * Inches(2.9)
    add_card(s2, cx, Inches(2.0), card_w, card_h)
    box = s2.shapes.add_textbox(cx + Inches(0.2), Inches(2.2), card_w - Inches(0.4), card_h - Inches(0.4))
    tf = box.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(18)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_BROWN_DEEP
    
    p2 = tf.add_paragraph()
    p2.text = sub
    p2.font.name = FONT_BODY
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_ORANGE
    p2.space_after = Pt(14)
    
    p3 = tf.add_paragraph()
    p3.text = desc
    p3.font.name = FONT_BODY
    p3.font.size = Pt(13)
    p3.font.color.rgb = COLOR_BROWN_TEXT

# ==================== SLIDE 3 ====================
s3 = create_slide()
add_header(s3, "STEP 01 · FEEDBACK", "1회차 소중한 피드백을 반영했어요")

add_card(s3, Inches(1.0), Inches(2.0), Inches(5.45), Inches(4.5))
box_l = s3.shapes.add_textbox(Inches(1.3), Inches(2.3), Inches(4.85), Inches(3.9))
tfl = box_l.text_frame
tfl.word_wrap = True
p = tfl.paragraphs[0]
p.text = "💬 보내주신 의견"
p.font.name = FONT_TITLE
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = COLOR_BROWN_DEEP
p.space_after = Pt(14)

fb_list = [
    "• 워크북에 적을 내용이 많아 살짝 부담스러웠어요.",
    "• 소감 나눔 시간이 조금 길게 느껴졌어요.",
    "• 워크북(이론)과 바이브 코딩(실습)이 더 유기적으로 연결되면 좋겠어요!",
    "• 내 개인 상황에 꼭 맞는 루틴도 짜보고 싶어요."
]
for item in fb_list:
    pi = tfl.add_paragraph()
    pi.text = item
    pi.font.name = FONT_BODY
    pi.font.size = Pt(14)
    pi.font.color.rgb = COLOR_BROWN_TEXT
    pi.space_after = Pt(10)

add_card(s3, Inches(6.883), Inches(2.0), Inches(5.45), Inches(4.5), COLOR_PEACH_LIGHT, COLOR_ORANGE)
box_r = s3.shapes.add_textbox(Inches(7.183), Inches(2.3), Inches(4.85), Inches(3.9))
tfr = box_r.text_frame
tfr.word_wrap = True
p = tfr.paragraphs[0]
p.text = "✨ 2회차의 약속 & 업그레이드"
p.font.name = FONT_TITLE
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = COLOR_BROWN_DEEP
p.space_after = Pt(14)

promise_list = [
    "✅ 복잡한 작성 부담은 덜고, 직관적인 실습 위주로!",
    "✅ 소감 나눔은 핵심만 콤팩트하게 (10분 컷)",
    "✅ 1회차에 정한 '나의 딥워크'를 화면 속 타이머로 바로 구현!",
    "✅ 소중한 목소리에 귀 기울여 계속 발전시키겠습니다 🌿"
]
for item in promise_list:
    pi = tfr.add_paragraph()
    pi.text = item
    pi.font.name = FONT_BODY
    pi.font.size = Pt(14)
    pi.font.bold = True
    pi.font.color.rgb = COLOR_BROWN_DEEP
    pi.space_after = Pt(10)

# ==================== SLIDE 4 ====================
s4 = create_slide()
add_header(s4, "STEP 01 · SHARING", "Q. 지난 시간 이후, 어떠셨나요?")

add_card(s4, Inches(1.0), Inches(2.0), Inches(11.333), Inches(4.8))
box = s4.shapes.add_textbox(Inches(1.5), Inches(2.4), Inches(10.333), Inches(4.0))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "일상에서 나의 시간 사용 & 시간 도둑을 관찰해보셨나요?"
p.font.name = FONT_TITLE
p.font.size = Pt(22)
p.font.color.rgb = COLOR_BROWN_DEEP
p.space_after = Pt(18)

p1 = tf.add_paragraph()
p1.text = "💡 수강생 나눔 예시 (채팅창에 한 줄로 편하게 남겨주세요!)"
p1.font.name = FONT_BODY
p1.font.size = Pt(15)
p1.font.bold = True
p1.font.color.rgb = COLOR_ORANGE
p1.space_after = Pt(10)

exs = [
    '“저는 하루 4시간 딥워크를 시도해봤는데, 마음먹으니 2~3시간은 몰입이 가능하더라고요!”',
    '“생각보다 시간 도둑이 너무 많아서 작업 시작 자체가 어렵다는 걸 깨달았어요 ㅠㅠ”',
    '“바다/언덕/평원을 구분해두니까 죄책감 없이 쉴 수 있어서 좋았습니다.”'
]
for ex in exs:
    pe = tf.add_paragraph()
    pe.text = ex
    pe.font.name = FONT_BODY
    pe.font.size = Pt(14)
    pe.font.color.rgb = COLOR_BROWN_TEXT
    pe.space_after = Pt(8)

# ==================== SLIDE 5 ====================
s5 = create_slide()
add_header(s5, "STEP 01 · BEHIND", "원래 저희의 야심찬(?) 의도는...")

add_card(s5, Inches(1.5), Inches(2.2), Inches(10.333), Inches(4.4))
box = s5.shapes.add_textbox(Inches(2.0), Inches(2.6), Inches(9.333), Inches(3.6))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "🔥 원래 상상했던 단톡방 풍경"
p.font.name = FONT_TITLE
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = COLOR_ORANGE
p.space_after = Pt(10)

p1 = tf.add_paragraph()
p1.text = "“다 함께 으쌰으쌰 인증샷 올리고, 루틴 자랑하고, 서로 불태우는 뜨거운 모임!”"
p1.font.name = FONT_BODY
p1.font.size = Pt(16)
p1.font.color.rgb = COLOR_BROWN_TEXT
p1.space_after = Pt(24)

p2 = tf.add_paragraph()
p2.text = " 현실의 단톡방 : ... (고요) 👀"
p2.font.name = FONT_TITLE
p2.font.size = Pt(20)
p2.font.bold = True
p2.font.color.rgb = COLOR_BROWN_DEEP
p2.space_after = Pt(10)

p3 = tf.add_paragraph()
p3.text = "“...제가 더 열심히 노력하겠습니다!! ㅋㅋㅋ 오늘 타이머 만들고 나면 자랑할 거 엄청 많아질 거예요!”"
p3.font.name = FONT_BODY
p3.font.size = Pt(16)
p3.font.color.rgb = COLOR_ORANGE

# ==================== SLIDE 6 ====================
s6 = create_slide()
add_header(s6, "STEP 01 · AI CHECK", "Q. 나의 AI 사용 경험은 어떠신가요?")

add_card(s6, Inches(1.0), Inches(2.0), Inches(5.45), Inches(4.5))
box_l = s6.shapes.add_textbox(Inches(1.3), Inches(2.4), Inches(4.85), Inches(3.7))
tfl = box_l.text_frame
tfl.word_wrap = True
p = tfl.paragraphs[0]
p.text = "💬 질문 2가지"
p.font.name = FONT_TITLE
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = COLOR_BROWN_DEEP
p.space_after = Pt(16)

p1 = tfl.add_paragraph()
p1.text = "1. 평소에 AI를 언제, 무슨 용도로 써보셨나요?\n\n2. 혹시 AI로 웹사이트나 웹페이지(HTML)를 만들어보신 경험이 있나요?"
p1.font.name = FONT_BODY
p1.font.size = Pt(15)
p1.font.color.rgb = COLOR_BROWN_TEXT

add_card(s6, Inches(6.883), Inches(2.0), Inches(5.45), Inches(4.5), COLOR_PEACH_LIGHT, COLOR_ORANGE)
box_r = s6.shapes.add_textbox(Inches(7.183), Inches(2.4), Inches(4.85), Inches(3.7))
tfr = box_r.text_frame
tfr.word_wrap = True
p = tfr.paragraphs[0]
p.text = "💡 수강생 반응 예시"
p.font.name = FONT_TITLE
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = COLOR_BROWN_DEEP
p.space_after = Pt(16)

p2 = tfr.add_paragraph()
p2.text = "“그림 그려줘! 하고 그냥 장난치며 놀아본 정도예요 ㅎㅎ”\n\n“여행 갈 때 일정표 만들어달라고 해본 적 있어요!”\n\n“코딩은 아예 해본 적 없어서 오늘 엄청 떨려요...”"
p2.font.name = FONT_BODY
p2.font.size = Pt(14)
p2.font.color.rgb = COLOR_BROWN_TEXT

# ==================== SLIDE 7 ====================
s7 = create_slide()
add_card(s7, Inches(2.0), Inches(1.8), Inches(9.333), Inches(4.2), COLOR_PEACH_LIGHT, COLOR_ORANGE)
box = s7.shapes.add_textbox(Inches(2.5), Inches(2.5), Inches(8.333), Inches(2.8))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "걱정하지 마세요! 🌿"
p.alignment = PP_ALIGN.CENTER
p.font.name = FONT_TITLE
p.font.size = Pt(36)
p.font.bold = True
p.font.color.rgb = COLOR_BROWN_DEEP
p.space_after = Pt(16)

p1 = tf.add_paragraph()
p1.text = "코딩 한 줄 몰라도 괜찮습니다.\n깨굴이와 꿀차가 차근차근 함께 만들어갈게요!"
p1.alignment = PP_ALIGN.CENTER
p1.font.name = FONT_BODY
p1.font.size = Pt(20)
p1.font.color.rgb = COLOR_BROWN_TEXT

# ==================== SLIDE 8 ====================
s8 = create_slide()
add_header(s8, "STEP 02 · TODAY'S GOAL", "오늘 함께 만들 실습 예제")

add_card(s8, Inches(1.5), Inches(2.0), Inches(10.333), Inches(4.8))
box = s8.shapes.add_textbox(Inches(2.0), Inches(2.4), Inches(9.333), Inches(4.0))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "🧗 나의 몰입 시간을 지켜줄 [딥워크 타이머]"
p.font.name = FONT_TITLE
p.font.size = Pt(26)
p.font.bold = True
p.font.color.rgb = COLOR_BROWN_DEEP
p.space_after = Pt(16)

features = [
    "⛰️ 산 오르기 : 타이머를 켜면 귀여운 캐릭터가 산을 등반하는 메인 화면",
    "🎯 딥워크 목표 : 이번 주 달성할 목표를 선택하고 연결하는 화면",
    "🧺 바구니 : 불쑥불쑥 떠오르는 잡생각/잡무를 던져두는 체크리스트",
    "⏱️ 업무시간 기록 : 오늘 내가 산을 얼마나 올랐는지 기록이 착착 쌓이는 화면"
]
for feat in features:
    pf = tf.add_paragraph()
    pf.text = feat
    pf.font.name = FONT_BODY
    pf.font.size = Pt(15)
    pf.font.color.rgb = COLOR_BROWN_TEXT
    pf.space_after = Pt(10)

# ==================== SLIDE 9 ====================
s9 = create_slide()
add_header(s9, "STEP 02 · STORY", "이 타이머를 만드는 데 얼마나 걸렸을까요?")

add_card(s9, Inches(2.0), Inches(2.2), Inches(9.333), Inches(4.4))
box = s9.shapes.add_textbox(Inches(2.5), Inches(2.8), Inches(8.333), Inches(3.2))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "👀 꿀차의 생생한 목격담"
p.alignment = PP_ALIGN.CENTER
p.font.name = FONT_TITLE
p.font.size = Pt(24)
p.font.bold = True
p.font.color.rgb = COLOR_ORANGE
p.space_after = Pt(20)

p1 = tf.add_paragraph()
p1.text = "클로드와 대화하며 뚝딱뚝딱 만드는 깨굴이...\n그리고 그 옆에서 눈을 껌뻑이며 구경하던 꿀차...\n\n과연 개발자 깨굴이는 이 페이지를 만드는 데 몇 분이 걸렸을까요?"
p1.alignment = PP_ALIGN.CENTER
p1.font.name = FONT_BODY
p1.font.size = Pt(18)
p1.font.color.rgb = COLOR_BROWN_TEXT

# ==================== SLIDE 10 ====================
s10 = create_slide()
add_card(s10, Inches(2.0), Inches(1.8), Inches(9.333), Inches(4.2), COLOR_PEACH_LIGHT, COLOR_ORANGE)
box = s10.shapes.add_textbox(Inches(2.5), Inches(2.5), Inches(8.333), Inches(2.8))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "놀랍게도 단 '5분'이면 충분했습니다!"
p.alignment = PP_ALIGN.CENTER
p.font.name = FONT_TITLE
p.font.size = Pt(32)
p.font.bold = True
p.font.color.rgb = COLOR_BROWN_DEEP
p.space_after = Pt(14)

p1 = tf.add_paragraph()
p1.text = "😲 꿀차 : “...네? 벌써 끝났다고요? 진짜로요?”\n(옆에서 턱 빠지게 놀라는 꿀차)"
p1.alignment = PP_ALIGN.CENTER
p1.font.name = FONT_BODY
p1.font.size = Pt(18)
p1.font.color.rgb = COLOR_ORANGE

# ==================== SLIDE 11 ====================
s11 = create_slide()
add_header(s11, "STEP 02 · SECRET", "어떻게 5분 만에 만들 수 있었을까요?")

add_card(s11, Inches(1.5), Inches(2.0), Inches(10.333), Inches(4.8))
box = s11.shapes.add_textbox(Inches(2.0), Inches(2.4), Inches(9.333), Inches(4.0))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "💡 핵심은 '의도'와 '기능'을 AI에게 명확하게 전달한 것!"
p.font.name = FONT_TITLE
p.font.size = Pt(22)
p.font.bold = True
p.font.color.rgb = COLOR_BROWN_DEEP
p.space_after = Pt(16)

secrets = [
    "1. 내가 무엇을 만들고 싶은지 명확한 기획 (산 오르기 컨셉, 4개 탭 구조)",
    "2. 각 버튼과 타이머가 어떻게 동작해야 하는지 명확한 규칙 지시",
    "3. AI와 단 3번의 티키타카 대화로 완성!",
    "\n(솔직한 후기: 사실 5분 만에 뼈대 나오고, 디테일 수정까지 딱 10분 걸렸습니다 ㅎㅎ)"
]
for sec in secrets:
    ps = tf.add_paragraph()
    ps.text = sec
    ps.font.name = FONT_BODY
    ps.font.size = Pt(15)
    ps.font.color.rgb = COLOR_BROWN_TEXT
    ps.space_after = Pt(8)

# ==================== SLIDE 12 ====================
s12 = create_slide()
add_card(s12, Inches(2.0), Inches(1.8), Inches(9.333), Inches(4.2))
box = s12.shapes.add_textbox(Inches(2.5), Inches(2.5), Inches(8.333), Inches(2.8))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "당황하지 마세요! 😌"
p.alignment = PP_ALIGN.CENTER
p.font.name = FONT_TITLE
p.font.size = Pt(32)
p.font.bold = True
p.font.color.rgb = COLOR_BROWN_DEEP
p.space_after = Pt(14)

p1 = tf.add_paragraph()
p1.text = "처음 AI에게 코딩을 시키면 생각과 다르게 나올 수 있어요.\n(버튼이 이상한 데 있거나 엉뚱한 이름이 붙거나...)\n\n하지만 괜찮습니다. 대화하며 하나씩 고쳐나가면 됩니다!"
p1.alignment = PP_ALIGN.CENTER
p1.font.name = FONT_BODY
p1.font.size = Pt(18)
p1.font.color.rgb = COLOR_BROWN_TEXT

# ==================== SLIDE 13 ====================
s13 = create_slide()
add_header(s13, "STEP 02 · VIBE CODING", "웹페이지를 만드는 3가지 핵심 요소")

w3 = Inches(3.55)
h3 = Inches(4.5)
data13 = [
    ("1. 기획하기 🧭", "내가 만들고 싶은 것은?", "• 의도와 쓰임새\n• 화면 구성과 레이아웃\n• 탭 이름과 버튼 용어"),
    ("2. 기능 만들기 ⚙️", "어떻게 작동해야 할까?", "• 타이머 시작 / 정지\n• 체크리스트 토글\n• 데이터 자동 저장 (localStorage)"),
    ("3. 디자인 입히기 🎨", "보면 기분이 좋아지게!", "• 밤하늘 & 오렌지 무드\n• 등산 캐릭터 애니메이션\n• 정갈한 폰트와 감성")
]
for i, (t, sub, desc) in enumerate(data13):
    cx = Inches(1.0) + i * Inches(3.89)
    add_card(s13, cx, Inches(2.0), w3, h3)
    b = s13.shapes.add_textbox(cx + Inches(0.25), Inches(2.3), w3 - Inches(0.5), h3 - Inches(0.6))
    tf = b.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = t
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(20)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_BROWN_DEEP
    
    p2 = tf.add_paragraph()
    p2.text = sub
    p2.font.name = FONT_BODY
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_ORANGE
    p2.space_after = Pt(14)
    
    p3 = tf.add_paragraph()
    p3.text = desc
    p3.font.name = FONT_BODY
    p3.font.size = Pt(14)
    p3.font.color.rgb = COLOR_BROWN_TEXT

# ==================== SLIDE 14 ====================
s14 = create_slide()
add_header(s14, "STEP 02 · FOCUS", "오늘 2회차 모임에서 집중할 것!")

add_card(s14, Inches(1.5), Inches(2.0), Inches(10.333), Inches(4.8), COLOR_PEACH_LIGHT, COLOR_ORANGE)
box = s14.shapes.add_textbox(Inches(2.0), Inches(2.4), Inches(9.333), Inches(4.0))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "🎯 오늘의 목표 : '기획'과 '기능'에 먼저 집중해요!"
p.font.name = FONT_TITLE
p.font.size = Pt(24)
p.font.bold = True
p.font.color.rgb = COLOR_BROWN_DEEP
p.space_after = Pt(18)

items14 = [
    ("✅ 1. 기획", "내가 쓰고 싶은 의도와 탭 구성 명확히 잡기 (오늘 완성!)"),
    ("✅ 2. 기능", "클릭, 타이머 작동, 새로고침해도 안 날아가는 저장 구현 (오늘 완성!)"),
    ("⏳ 3. 디자인", "디자인은 깨굴이가 준비한 '치트키 템플릿'으로 깔끔하게 맞추고, 세부 커스텀은 3회차에서 더 깊게 다뤄요!")
]
for title, desc in items14:
    p_t = tf.add_paragraph()
    p_t.text = title
    p_t.font.name = FONT_BODY
    p_t.font.size = Pt(16)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_ORANGE
    
    p_d = tf.add_paragraph()
    p_d.text = desc
    p_d.font.name = FONT_BODY
    p_d.font.size = Pt(14)
    p_d.font.color.rgb = COLOR_BROWN_TEXT
    p_d.space_after = Pt(12)

# ==================== SLIDE 15 ====================
s15 = create_slide()
add_header(s15, "STEP 03 · PLANNING", "[기획하기] 딥워크 타이머의 큰 그림")

add_card(s15, Inches(1.5), Inches(2.0), Inches(10.333), Inches(4.8))
box = s15.shapes.add_textbox(Inches(2.0), Inches(2.4), Inches(9.333), Inches(4.0))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "내가 진짜로 쓰고 싶은 도구는 무엇인가요?"
p.font.name = FONT_TITLE
p.font.size = Pt(22)
p.font.bold = True
p.font.color.rgb = COLOR_BROWN_DEEP
p.space_after = Pt(16)

p1 = tf.add_paragraph()
p1.text = "🧗 예제 컨셉 : 시간의 흐름에 따라 산을 오르는 '딥워크 등반 타이머'"
p1.font.name = FONT_BODY
p1.font.size = Pt(16)
p1.font.bold = True
p1.font.color.rgb = COLOR_ORANGE
p1.space_after = Pt(14)

desc15 = [
    "• 단순히 숫자만 줄어드는 타이머는 지루하고 압박감만 줍니다.",
    "• '내가 오늘 4시간의 산을 몇 % 정복했는지' 눈으로 직접 확인하는 시각적 성취감!",
    "• 잡생각은 바구니에 던져두고, 오직 지금 오르는 봉우리에만 몰입할 수 있는 환경을 만듭니다."
]
for d in desc15:
    pd = tf.add_paragraph()
    pd.text = d
    pd.font.name = FONT_BODY
    pd.font.size = Pt(14)
    pd.font.color.rgb = COLOR_BROWN_TEXT
    pd.space_after = Pt(8)

# ==================== SLIDE 16 ====================
s16 = create_slide()
add_header(s16, "STEP 03 · ACTIVITY", "[기획하기] 각 탭의 핵심 목적 한 줄 정의")

add_card(s16, Inches(1.0), Inches(2.0), Inches(11.333), Inches(4.8))
box = s16.shapes.add_textbox(Inches(1.5), Inches(2.3), Inches(10.333), Inches(4.2))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "✍️ 4가지 탭이 나에게 왜 필요한지 '한 줄'로 적어보아요!"
p.font.name = FONT_TITLE
p.font.size = Pt(22)
p.font.bold = True
p.font.color.rgb = COLOR_BROWN_DEEP
p.space_after = Pt(14)

p_guide = tf.add_paragraph()
p_guide.text = "완벽한 정답은 없습니다. 내가 생각하는 쓰임새를 자유롭게 채팅창에 올려주세요!"
p_guide.font.name = FONT_BODY
p_guide.font.size = Pt(14)
p_guide.font.color.rgb = COLOR_ORANGE
p_guide.space_after = Pt(16)

tabs16 = [
    "• ⛰️ 산 오르기 : ________________________________________________",
    "• 🎯 딥워크 : __________________________________________________",
    "• 🧺 바구니 : __________________________________________________",
    "• ⏱️ 업무시간 : _________________________________________________"
]
for t in tabs16:
    pt = tf.add_paragraph()
    pt.text = t
    pt.font.name = FONT_BODY
    pt.font.size = Pt(15)
    pt.font.color.rgb = COLOR_BROWN_TEXT
    pt.space_after = Pt(10)

# ==================== SLIDE 17 ====================
s17 = create_slide()
add_header(s17, "STEP 03 · REFERENCE", "[기획하기] 4개 탭의 기획 예시")

w4 = Inches(2.6)
h4 = Inches(4.5)
data17 = [
    ("⛰️ 산 오르기", "오늘의 작업량 측정", "하루 목표(4시간) 대비\n현재 내가 얼마나 올라왔는지\n캐릭터와 게이지로 확인하는 탭"),
    ("🎯 딥워크", "몰입할 일거리 지정", "이번 주 목표 중에서\n오늘 집중해서 오를 일거리를\n선택하고 설정하는 탭"),
    ("🧺 바구니", "잡생각 임시 보관소", "딥워크 중에 불쑥 떠오르는\n딴짓거리나 잔업을 마감 요일과\n함께 툭 던져두는 체크리스트"),
    ("⏱️ 업무시간", "나의 발자취 기록", "타이머를 멈출 때마다\n오늘과 이번 주에 얼마나 올랐는지\n시간 기록이 착착 쌓이는 탭")
]
for i, (t, sub, desc) in enumerate(data17):
    cx = Inches(1.0) + i * Inches(2.9)
    add_card(s17, cx, Inches(2.0), w4, h4)
    b = s17.shapes.add_textbox(cx + Inches(0.2), Inches(2.2), w4 - Inches(0.4), h4 - Inches(0.4))
    tf = b.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = t
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(18)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_BROWN_DEEP
    
    p2 = tf.add_paragraph()
    p2.text = sub
    p2.font.name = FONT_BODY
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_ORANGE
    p2.space_after = Pt(12)
    
    p3 = tf.add_paragraph()
    p3.text = desc
    p3.font.name = FONT_BODY
    p3.font.size = Pt(13)
    p3.font.color.rgb = COLOR_BROWN_TEXT

# ==================== SLIDE 18 ====================
s18 = create_slide()
add_header(s18, "STEP 04 · LIVE DEMO", "이제 시연이의 실시간 코딩 시연이 시작됩니다!")

add_card(s18, Inches(1.5), Inches(2.0), Inches(10.333), Inches(4.8), COLOR_PEACH_LIGHT, COLOR_ORANGE)
box = s18.shapes.add_textbox(Inches(2.0), Inches(2.5), Inches(9.333), Inches(3.8))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "🐸 개발자 깨굴이(이시연)의 바이브 코딩 LIVE"
p.font.name = FONT_TITLE
p.font.size = Pt(24)
p.font.bold = True
p.font.color.rgb = COLOR_BROWN_DEEP
p.space_after = Pt(16)

demos = [
    "1. 뼈대 프롬프트 입력 : 기획 내용과 4개 탭 레이아웃 한 번에 잡기",
    "2. 텍스트 와이어프레임 점검 : 엉뚱한 버튼명이나 깨진 기능 바로잡기",
    "3. 깨굴이의 '치트키 디자인 코드' 주입 : 밤하늘 감성과 등산 애니메이션 완성!",
    "4. 여러분도 곧이어 직접 만들어볼 거예요 (막히는 부분은 1:1로 봐드립니다!)"
]
for d in demos:
    pd = tf.add_paragraph()
    pd.text = d
    pd.font.name = FONT_BODY
    pd.font.size = Pt(15)
    pd.font.color.rgb = COLOR_BROWN_TEXT
    pd.space_after = Pt(10)

# ==================== SLIDE 19 ====================
s19 = create_slide()
add_header(s19, "PREVIEW · 3회차", "3회차 모임 미리보기")

add_card(s19, Inches(1.0), Inches(2.0), Inches(5.45), Inches(4.5))
box_l = s19.shapes.add_textbox(Inches(1.3), Inches(2.4), Inches(4.85), Inches(3.7))
tfl = box_l.text_frame
tfl.word_wrap = True
p = tfl.paragraphs[0]
p.text = "🌐 깃허브(GitHub)로 무료 배포하기"
p.font.name = FONT_TITLE
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = COLOR_BROWN_DEEP
p.space_after = Pt(14)

p1 = tfl.add_paragraph()
p1.text = "• 컴퓨터 파일(HTML)로만 열리던 내 타이머를 진짜 웹사이트로!\n\n• 스마트폰 홈 화면에 아이콘으로 추가해서 매일 앱처럼 사용하는 법을 배워요."
p1.font.name = FONT_BODY
p1.font.size = Pt(14)
p1.font.color.rgb = COLOR_BROWN_TEXT

add_card(s19, Inches(6.883), Inches(2.0), Inches(5.45), Inches(4.5))
box_r = s19.shapes.add_textbox(Inches(7.183), Inches(2.4), Inches(4.85), Inches(3.7))
tfr = box_r.text_frame
tfr.word_wrap = True
p = tfr.paragraphs[0]
p.text = "🛠️ 나만의 맞춤 기능 & 디자인 디벨롭"
p.font.name = FONT_TITLE
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = COLOR_BROWN_DEEP
p.space_after = Pt(14)

p2 = tfr.add_paragraph()
p2.text = "• 등산 말고 수영, 우주비행 등 나만의 컨셉으로 커스텀!\n\n• 직접 써보며 느낀 불편한 점을 고치고, 나에게 꼭 필요한 추가 기능을 장착해 봅니다."
p2.font.name = FONT_BODY
p2.font.size = Pt(14)
p2.font.color.rgb = COLOR_BROWN_TEXT

# ==================== SLIDE 20 ====================
s20 = create_slide()
add_header(s20, "MISSION · 2회차 숙제", "이번 주 실천 미션 & 숙제")

add_card(s20, Inches(1.5), Inches(2.0), Inches(10.333), Inches(4.8), COLOR_PEACH_LIGHT, COLOR_ORANGE)
box = s20.shapes.add_textbox(Inches(2.0), Inches(2.4), Inches(9.333), Inches(4.0))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "🌿 이번 주 실천 약속 2가지"
p.font.name = FONT_TITLE
p.font.size = Pt(24)
p.font.bold = True
p.font.color.rgb = COLOR_BROWN_DEEP
p.space_after = Pt(16)

hw_list = [
    ("1. 나의 딥워크 시간 '인증샷' 올려보기 (선택)", "오늘 만든 타이머를 직접 켜고 딥워크를 실천한 뒤 캡처해서 올려보세요!"),
    ("2. 내가 만든 타이머 직접 써보고 아쉬운 점 적어두기", "“이 버튼이 여기에 있으면 좋겠네”, “이 기능이 추가되면 편하겠다” 메모해오기!"),
    ("💌 오픈채팅방은 언제나 열려있어요!", "만들다 막히는 부분은 언제든 옾챗에 질문 주시면 깨굴이와 꿀차가 봐드립니다. 결과물 자랑도 대환영이에요! ✨")
]
for title, desc in hw_list:
    p_t = tf.add_paragraph()
    p_t.text = title
    p_t.font.name = FONT_BODY
    p_t.font.size = Pt(16)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_ORANGE
    
    p_d = tf.add_paragraph()
    p_d.text = desc
    p_d.font.name = FONT_BODY
    p_d.font.size = Pt(14)
    p_d.font.color.rgb = COLOR_BROWN_TEXT
    p_d.space_after = Pt(10)

# Save
output_path = '/Users/hnky/agyspace/mytt_2회차_강의슬라이드_구원파트.pptx'
prs.save(output_path)
print("Successfully generated:", output_path, "Slides count:", len(prs.slides))
