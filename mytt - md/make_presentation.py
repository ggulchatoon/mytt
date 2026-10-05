import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# 1. Colors & Constants
FONT_TITLE = 'Gowun Dodum'
FONT_BODY = 'Noto Sans KR'

COLOR_BROWN_DEEP = RGBColor(62, 44, 33)      # #3E2C21 (제목 짙은 브라운)
COLOR_BROWN_TEXT = RGBColor(92, 74, 62)      # #5C4A3E (본문 텍스트)
COLOR_BROWN_MUTED = RGBColor(140, 115, 95)   # #8C735F (설명/힌트)
COLOR_ORANGE = RGBColor(240, 164, 92)        # #F0A45C (포인트 오렌지)
COLOR_PEACH_CARD = RGBColor(255, 248, 240)   # #FFF8F0 (카드 배경 크림)
COLOR_PEACH_BORDER = RGBColor(237, 216, 200) # #EDD8C8 (카드 테두리)
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_DARK_PANEL = RGBColor(27, 32, 48)      # 다크 테마 카드

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

bg_image_path = '/Users/hnky/agyspace/slide_bg.png'

def set_slide_background(slide):
    if os.path.exists(bg_image_path):
        slide.shapes.add_picture(bg_image_path, 0, 0, prs.slide_width, prs.slide_height)

def add_header(slide, step_text="mytt 2회차 모임", main_title=""):
    # Step label
    if step_text:
        step_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.333), Inches(0.4))
        tf_s = step_box.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0
        p_s = tf_s.paragraphs[0]
        p_s.text = step_text.upper()
        p_s.font.name = FONT_BODY
        p_s.font.size = Pt(13)
        p_s.font.bold = True
        p_s.font.color.rgb = COLOR_ORANGE

    # Main Title
    if main_title:
        title_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.0), Inches(11.333), Inches(0.8))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = main_title
        p_t.font.name = FONT_TITLE
        p_t.font.size = Pt(28)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_BROWN_DEEP

def add_card(slide, left, top, width, height, bg_color=COLOR_PEACH_CARD, border_color=COLOR_PEACH_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape

print("Script template ready.")
